import json
import csv
import re
import time
import tempfile
from contextlib import contextmanager
import requests
from pathlib import Path
from datetime import datetime, timezone
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Font, Alignment
from textual import work
from textual.app import App, ComposeResult
from textual.screen import ModalScreen, Screen
from textual.containers import Vertical, Horizontal
from textual.widgets import (
    Header, Footer, DataTable, TabbedContent, TabPane,
    Input, Button, Label, Static
)

OUTPUT_DIR = "output"
CACHE_DIR = "cache"


@contextmanager
def report_file(final_path: Path):
    cache_dir = Path(__file__).resolve().parent / CACHE_DIR
    cache_dir.mkdir(exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=cache_dir, prefix=".report-", suffix=final_path.suffix, delete=False) as file:
        working_path = Path(file.name)
    try:
        yield working_path
        working_path.replace(final_path)
    finally:
        working_path.unlink(missing_ok=True)

API_HEADERS = {
    "User-Agent":      "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
                       "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
    "Accept":          "application/json, text/plain, */*",
    "Accept-Language": "en-US,en;q=0.9",
    "Referer":         "https://www.exophase.com/",
}

AWARDS_COLUMNS = ("Achievement", "Earned (UTC)", "Earned (site)", "Award ID")
HISTORY_COLUMNS = ("Game", "Platform", "Achievement", "Earned (UTC)", "Earned (site)", "Award ID")

PLATFORM_RULES = [
    ("PlayStation", ["playstation", "ps4", "ps5", "ps3", "ps vita"]),
    ("Xbox",        ["xbox"]),
    ("Steam",       ["steam"]),
]

def classify_platform(platforms: list[dict]) -> str:
    names_lower = [str(p.get("name", "")).lower() for p in platforms if isinstance(p, dict)]
    for tab_name, keywords in PLATFORM_RULES:
        if any(kw in name for name in names_lower for kw in keywords):
            return tab_name
    return "Other"

def fmt_playtime(game: dict) -> str:
    units = game.get("playtimeUnits") or {}
    h, m = units.get("hours", 0), units.get("minutes", 0)
    return f"{h}h {m}m" if (h or m) else str(game.get("playtime", "0h"))

def fmt_timestamp(ts) -> str:
    if ts and ts > 0:
        try:
            return datetime.fromtimestamp(ts, timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
        except Exception:
            pass
    return ""

def build_row(game: dict, platform: str) -> tuple:
    meta       = game.get("meta") or {}
    title      = meta.get("title", "Unknown")
    playtime   = fmt_playtime(game)
    percent    = game.get("percent", 0.0)
    lastplayed = fmt_timestamp(game.get("lastplayed_utc", 0))

    if platform == "PlayStation":
        return (
            title, playtime,
            game.get("earned_bronze", 0) or 0,
            game.get("earned_silver", 0) or 0,
            game.get("earned_gold", 0) or 0,
            game.get("earned_platinum", 0) or 0,
            percent, lastplayed,
        )
    if platform == "Xbox":
        return (
            title, playtime,
            game.get("earned_awards", 0) or 0,
            game.get("total_awards", 0) or 0,
            game.get("earned_points", 0) or 0,
            percent, lastplayed,
        )
    if platform == "Steam":
        return (
            title, playtime,
            game.get("earned_awards", 0) or 0,
            game.get("total_awards", 0) or 0,
            percent, lastplayed,
        )
    platforms_str = ", ".join(
        str(p.get("name", "")) for p in (meta.get("platforms") or []) if isinstance(p, dict)
    )
    return (title, playtime, platforms_str, percent, lastplayed)


def pretty_award_name(slug: str) -> str:
    if not slug:
        return "Unknown"
    s = str(slug).replace("-", " ").replace("_", " ").strip()
    return s.title() if s else "Unknown"


def get_game_ids(game: dict) -> tuple:
    """Повертає (master_playerid, master_id) для запиту /earned.

    games-API віддає обидва поля на верхньому рівні (master_playerid,
    master_id), meta.master_id дублює master_id. master_playerid — це
    environment player id (напр. Steam id), а НЕ legacy playerProfileId,
    тому для /earned треба брати саме його.
    """
    if not isinstance(game, dict):
        return (None, None)
    meta = game.get("meta") or {}
    pid = game.get("master_playerid") or meta.get("master_playerid")
    gid = game.get("master_id") or meta.get("master_id")
    # fallback: деякі старі дампи можуть мати інші ключі
    if pid is None:
        pid = game.get("masterPlayerId") or game.get("player_id")
    if gid is None:
        gid = game.get("game_id") or game.get("id")
    try:
        pid = int(pid) if pid is not None else None
    except (ValueError, TypeError):
        pass
    try:
        gid = int(gid) if gid is not None else None
    except (ValueError, TypeError):
        pass
    return (pid, gid)


def build_award_row(award: dict) -> tuple:
    slug = award.get("slug", "")
    name = pretty_award_name(slug)
    ts = award.get("timestamp", 0) or 0
    earned_utc = fmt_timestamp(ts)
    earned_site = str(award.get("earned", "") or "")
    aid = award.get("awardid") or award.get("masterAwardId") or ""
    return (name, earned_utc, earned_site, str(aid))


def awards_cache_path(cache_dir: Path, pid, gid) -> Path:
    return Path(cache_dir) / f"awards_{pid}_{gid}.json"


def fetch_earned_awards(master_playerid, master_id, timeout: int = 15) -> list[dict]:
    """Синхронно тягне https://api.exophase.com/public/player/{pid}/game/{gid}/earned.

    Успіх = {"list": [...]}, помилка (не той pid) = {"success": false}.
    """
    url = f"https://api.exophase.com/public/player/{master_playerid}/game/{master_id}/earned"
    resp = requests.get(url, headers=API_HEADERS, timeout=timeout)
    resp.raise_for_status()
    data = resp.json()
    lst = data.get("list")
    if lst is None:
        # {"success": false} або неочікуваний формат
        raise ValueError(f"API did not return an achievement list (pid={master_playerid} gid={master_id}): {str(data)[:200]}")
    return lst if isinstance(lst, list) else []


MODAL_CSS = """
ModalScreen {
    align: center middle;
    background: $background 80%;
}
#dialog {
    width: 64;
    height: auto;
    padding: 2 3;
    background: $surface;
    border: solid $accent;
}
#dialog Label { margin-bottom: 1; text-style: bold; }
#dialog Input  { margin-bottom: 1; }
Horizontal { margin-top: 1; height: auto; align: right middle; }
Button { margin-left: 1; }
"""

class SyncModal(ModalScreen[str | None]):
    CSS = MODAL_CSS

    def compose(self) -> ComposeResult:
        with Vertical(id="dialog"):
            yield Label("Synchronise with Exophase")
            yield Input(placeholder="Player ID", id="pid")
            with Horizontal():
                yield Button("Cancel", variant="error",   id="cancel")
                yield Button("Synchronise", variant="success", id="start")

    def on_button_pressed(self, event: Button.Pressed) -> None:
        if event.button.id == "start":
            pid = self.query_one("#pid", Input).value.strip()
            if pid:
                self.dismiss(pid)
            else:
                self.app.notify("Enter a Player ID.", severity="warning")
        else:
            self.dismiss(None)


class FilterModal(ModalScreen[str | None]):
    CSS = MODAL_CSS

    def compose(self) -> ComposeResult:
        with Vertical(id="dialog"):
            yield Label("Filter by game title")
            yield Input(placeholder="Enter part of a title…", id="query")
            with Horizontal():
                yield Button("Reset", variant="warning", id="reset")
                yield Button("Cancel", variant="error", id="cancel")
                yield Button("Apply", variant="success", id="apply")

    def on_button_pressed(self, event: Button.Pressed) -> None:
        if event.button.id == "apply":
            self.dismiss(self.query_one("#query", Input).value.strip())
        elif event.button.id == "reset":
            self.dismiss("")
        else:
            self.dismiss(None)


class AwardsScreen(Screen):
    """Екран списку ачивок конкретної гри: назва + дата отримання."""

    CSS = """
    DataTable { width: 1fr; height: 1fr; }
    #awards-title { height: 1; padding: 0 2; text-style: bold; }
    #awards-status { height: 1; background: $surface; color: $text-muted; padding: 0 2; }
    """

    BINDINGS = [
        ("q", "back", "Back"),
        ("escape", "back", "Back"),
        ("e", "export_excel", "Export Excel"),
        ("c", "export_csv", "Export CSV"),
        ("r", "refresh", "Refresh from API"),
    ]

    def __init__(self, game_title: str, platform_label: str,
                 master_playerid, master_id, out_dir: Path, cache_dir: Path) -> None:
        super().__init__()
        self.game_title = game_title
        self.platform_label = platform_label
        self.master_playerid = master_playerid
        self.master_id = master_id
        self.out_dir = Path(out_dir)
        self.cache_dir = Path(cache_dir)
        self.awards: list[dict] = []
        self.sort_state: dict[str, bool] = {}

    def compose(self) -> ComposeResult:
        yield Header(show_clock=True)
        yield Label(f"🏆 {self.game_title}  [{self.platform_label}]  (pid={self.master_playerid} gid={self.master_id})",
                    id="awards-title")
        yield DataTable(id="dt-awards")
        yield Static("Loading…", id="awards-status")
        yield Footer()

    def on_mount(self) -> None:
        dt = self.query_one("#dt-awards", DataTable)
        dt.cursor_type = "row"
        dt.zebra_stripes = True
        dt.add_columns(*AWARDS_COLUMNS)
        # 1) одразу показати кеш, якщо є
        try:
            cp = awards_cache_path(self.cache_dir, self.master_playerid, self.master_id)
            if cp.exists():
                with open(cp, "r", encoding="utf-8") as f:
                    cached = json.load(f)
                if isinstance(cached, list) and cached:
                    self.awards = cached
                    self._show_awards(f"📁 cache: {cp.name}")
                    # тихо оновити у фоні
                    self.fetch_data()
                    return
        except Exception:
            pass
        self.fetch_data()

    def _show_awards(self, status: str = "") -> None:
        dt = self.query_one("#dt-awards", DataTable)
        dt.clear()
        # найновіші зверху
        rows = sorted(self.awards, key=lambda a: a.get("timestamp", 0) or 0, reverse=True)
        for aw in rows:
            if isinstance(aw, dict):
                dt.add_row(*build_award_row(aw))
        base = f"🏆 {len(rows)} achievements — {self.game_title}"
        if status:
            base += f"   |   {status}"
        try:
            self.query_one("#awards-status", Static).update(base)
        except Exception:
            pass

    def _set_status(self, msg: str) -> None:
        try:
            self.query_one("#awards-status", Static).update(msg)
        except Exception:
            pass

    @work(thread=True)
    def fetch_data(self) -> None:
        self.app.call_from_thread(self._set_status, "🌐 Loading from Exophase API…")
        try:
            lst = fetch_earned_awards(self.master_playerid, self.master_id)
        except Exception as e:
            self.app.call_from_thread(
                self.app.notify, f"Error loading achievements: {e}", severity="error"
            )
            if not self.awards:
                self.app.call_from_thread(self._set_status, f"❌ Error: {e}")
            return
        self.awards = lst
        # кеш
        try:
            self.cache_dir.mkdir(exist_ok=True)
            cp = awards_cache_path(self.cache_dir, self.master_playerid, self.master_id)
            with open(cp, "w", encoding="utf-8") as f:
                json.dump(lst, f, ensure_ascii=False, indent=2)
            status = f"✅ {len(lst)} achievements → {cp.name}"
        except Exception as e:
            status = f"✅ {len(lst)} achievements (cache not saved: {e})"
        self.app.call_from_thread(self._show_awards, status)
        self.app.call_from_thread(self.app.notify, status, severity="information")

    def action_back(self) -> None:
        self.app.pop_screen()

    def action_refresh(self) -> None:
        self.fetch_data()

    def on_data_table_header_selected(self, event: DataTable.HeaderSelected) -> None:
        dt = event.data_table
        col_key = event.column_key
        state_key = f"awards_{col_key}"
        reverse = self.sort_state.get(state_key, False)
        dt.sort(col_key, reverse=reverse)
        self.sort_state[state_key] = not reverse

    def _export_rows(self) -> list[dict]:
        rows = []
        for aw in sorted(self.awards, key=lambda a: a.get("timestamp", 0) or 0, reverse=True):
            if not isinstance(aw, dict):
                continue
            row = build_award_row(aw)
            rows.append(dict(zip(AWARDS_COLUMNS, row)))
        return rows

    def action_export_excel(self) -> None:
        rows = self._export_rows()
        if not rows:
            self.app.notify("No achievements to export", severity="warning")
            return
        safe = re.sub(r"[^\w\-]+", "_", self.game_title)[:50]
        path = self.out_dir / f"awards_{self.master_playerid}_{self.master_id}_{safe}.xlsx"
        try:
            df = pd.DataFrame(rows)
            with report_file(path) as working_path:
                with pd.ExcelWriter(working_path, engine="openpyxl") as writer:
                    df.to_excel(writer, sheet_name="Awards", index=False)
                    ws = writer.sheets["Awards"]
                    for cell in ws[1]:
                        cell.fill = PatternFill("solid", fgColor="1F4E79")
                        cell.font = Font(color="FFFFFF", bold=True)
                        cell.alignment = Alignment(horizontal="center")
                    for col_idx, col in enumerate(ws.columns, 1):
                        width = max((len(str(c.value)) for c in col if c.value), default=8)
                        ws.column_dimensions[get_column_letter(col_idx)].width = width + 2
            self.app.notify(f"✅ Excel → {path.name}", severity="information")
        except Exception as e:
            self.app.notify(f"Excel export error: {e}", severity="error")

    def action_export_csv(self) -> None:
        rows = self._export_rows()
        if not rows:
            self.app.notify("No achievements to export", severity="warning")
            return
        safe = re.sub(r"[^\w\-]+", "_", self.game_title)[:50]
        path = self.out_dir / f"awards_{self.master_playerid}_{self.master_id}_{safe}.csv"
        try:
            with report_file(path) as working_path:
                with open(working_path, "w", newline="", encoding="utf-8-sig") as f:
                    writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
                    writer.writeheader()
                    writer.writerows(rows)
            self.app.notify(f"✅ CSV → {path.name}", severity="information")
        except Exception as e:
            self.app.notify(f"CSV export error: {e}", severity="error")


TABLE_COLUMNS = {
    "dt-ps":    ("Game", "Playtime", "Bronze", "Silver", "Gold", "Platinum", "Completion %", "Last Played (UTC)"),
    "dt-xbox":  ("Game", "Playtime", "Earned Awards", "Total Awards", "Earned Points", "Completion %", "Last Played (UTC)"),
    "dt-steam": ("Game", "Playtime", "Earned Awards", "Total Awards", "Completion %", "Last Played (UTC)"),
    "dt-other": ("Game", "Playtime", "Platforms", "Completion %", "Last Played (UTC)"),
}

TAB_TO_DT = {
    "tp-ps":    "#dt-ps",
    "tp-xbox":  "#dt-xbox",
    "tp-steam": "#dt-steam",
    "tp-other": "#dt-other",
}

PLATFORM_TO_DT = {
    "PlayStation": "#dt-ps",
    "Xbox":        "#dt-xbox",
    "Steam":       "#dt-steam",
    "Other":       "#dt-other",
}


class SELFexophase(App):
    CSS = """
    DataTable { width: 1fr; height: 1fr; }
    #statusbar { height: 1; background: $surface; color: $text-muted; padding: 0 2; }
    """
    BINDINGS = [
        ("q", "quit",        "Quit"),
        ("s", "sync",        "Synchronise"),
        ("a", "open_awards", "Achievements"),
        ("enter", "open_awards", "Achievements"),
        ("w", "sync_all_awards", "Sync all achievements"),
        ("W", "sync_all_awards_force", "Resync achievements"),
        ("h", "export_awards_history", "Achievement history"),
        ("d", "delete_game", "Delete"),
        ("e", "export",      "Export Excel"),
        ("c", "export_csv",  "Export CSV"),
        ("f", "filter",      "Filter"),
        ("r", "reload",      "Reload"),
    ]

    def __init__(self):
        super().__init__()
        self.sort_state:      dict[str, bool] = {}
        self.games_data:      list[dict]      = []
        self.latest_json_path: Path | None    = None
        self._filter_query:   str             = ""
        self._awards_sync_running: bool = False
        self._awards_sync_stop:    bool = False

    def compose(self) -> ComposeResult:
        yield Header(show_clock=True)
        with TabbedContent():
            with TabPane("PlayStation", id="tp-ps"):
                yield DataTable(id="dt-ps")
            with TabPane("Xbox", id="tp-xbox"):
                yield DataTable(id="dt-xbox")
            with TabPane("Steam", id="tp-steam"):
                yield DataTable(id="dt-steam")
            with TabPane("Other", id="tp-other"):
                yield DataTable(id="dt-other")
        yield Static("", id="statusbar")
        yield Footer()

    def on_mount(self) -> None:
        for dt_id, cols in TABLE_COLUMNS.items():
            dt = self.query_one(f"#{dt_id}", DataTable)
            dt.cursor_type  = "row"
            dt.zebra_stripes = True
            dt.add_columns(*cols)
        self.load_data()

    def load_data(self, filter_query: str = "") -> None:
        for dt_id in TABLE_COLUMNS:
            self.query_one(f"#{dt_id}", DataTable).clear()

        files = list(self._cache_dir().glob("all_games_*.json"))
        if not files:
            self.games_data = []
            self._update_statusbar()
            return

        self.latest_json_path = max(files, key=lambda p: p.stat().st_mtime)
        try:
            with open(self.latest_json_path, "r", encoding="utf-8") as f:
                self.games_data = json.load(f)
        except Exception as e:
            self.notify(f"Error reading JSON: {e}", severity="error")
            return

        tables = {k: self.query_one(v, DataTable) for k, v in PLATFORM_TO_DT.items()}
        q = filter_query.lower()
        loaded = 0

        for game in self.games_data:
            if not isinstance(game, dict):
                continue
            meta  = game.get("meta") or {}
            title = meta.get("title", "Unknown")
            if q and q not in title.lower():
                continue

            platforms = meta.get("platforms") or []
            platform  = classify_platform(platforms)
            row       = build_row(game, platform)
            tables[platform].add_row(*row)
            loaded += 1

        indicator = f"  🔍 filter: «{filter_query}»" if filter_query else ""
        self._update_statusbar(loaded, indicator)

    def _update_statusbar(self, count: int = 0, extra: str = "") -> None:
        file_info = self.latest_json_path.name if self.latest_json_path else "—"
        self.query_one("#statusbar", Static).update(
            f" 📁 {file_info}   |   🎮 {count} games{extra}"
        )

    @staticmethod
    def _sort_key(value) -> tuple:
        if isinstance(value, (int, float)):
            return (0, float(value))
        if isinstance(value, str):
            if not value:
                return (0, 0.0)
            if "h " in value and "m" in value:
                try:
                    h, rest = value.split("h ")
                    m = rest.replace("m", "").strip()
                    return (0, int(h) * 60 + int(m))
                except ValueError:
                    pass
            try:
                return (0, float(value))
            except ValueError:
                return (1, value.lower())
        return (1, str(value).lower())

    def on_data_table_header_selected(self, event: DataTable.HeaderSelected) -> None:
        dt        = event.data_table
        col_key   = event.column_key
        state_key = f"{dt.id}_{col_key}"
        reverse   = self.sort_state.get(state_key, False)
        dt.sort(col_key, key=self._sort_key, reverse=reverse)
        self.sort_state[state_key] = not reverse

    def _active_datatable(self) -> DataTable | None:
        active = self.query_one(TabbedContent).active
        dt_selector = TAB_TO_DT.get(active)
        return self.query_one(dt_selector, DataTable) if dt_selector else None

    def _output_dir(self) -> Path:
        d = Path(__file__).resolve().parent / OUTPUT_DIR
        d.mkdir(exist_ok=True)
        return d

    def _cache_dir(self) -> Path:
        d = Path(__file__).resolve().parent / CACHE_DIR
        d.mkdir(exist_ok=True)
        return d

    def _save_games_data(self) -> None:
        if not self.latest_json_path:
            self.latest_json_path = self._cache_dir() / "all_games_manual.json"
        try:
            with open(self.latest_json_path, "w", encoding="utf-8") as f:
                json.dump(self.games_data, f, ensure_ascii=False, indent=4)
        except Exception as e:
            self.notify(f"Save error: {e}", severity="error")

    def action_reload(self) -> None:
        self.load_data(self._filter_query)
        self.notify("Data reloaded", severity="information")

    def action_filter(self) -> None:
        def apply_filter(query: str | None) -> None:
            if query is None:
                return
            self._filter_query = query
            self.load_data(query)
            if query:
                self.notify(f"Filter: «{query}»", severity="information")
            else:
                self.notify("Filter reset", severity="information")

        self.push_screen(FilterModal(), apply_filter)

    def _selected_game(self) -> tuple[dict | None, str]:
        """Повертає (game_dict, platform_label) для поточного рядка активної вкладки."""
        dt = self._active_datatable()
        if dt is None:
            return (None, "")
        try:
            row_key = dt.coordinate_to_cell_key(dt.cursor_coordinate).row_key
            row_values = dt.get_row(row_key)
            target_title = str(row_values[0])
        except Exception:
            return (None, "")
        try:
            tabs = self.query_one(TabbedContent)
            active_id = tabs.active
            platform = active_id.replace("tp-", "").capitalize()
            if platform == "Ps":
                platform = "PlayStation"
        except Exception:
            platform = ""
        # шукаємо гру за назвою + платформою вкладки (захист від дублів назв)
        fallback = None
        for game in self.games_data:
            if not isinstance(game, dict):
                continue
            meta = game.get("meta") or {}
            if meta.get("title", "") != target_title:
                continue
            plat = classify_platform(meta.get("platforms") or [])
            if plat.lower() == platform.lower():
                return (game, plat)
            if fallback is None:
                fallback = (game, plat)
        if fallback is not None:
            return fallback
        return (None, platform)

    def action_open_awards(self) -> None:
        game, platform = self._selected_game()
        if not game:
            self.notify("No row selected", severity="warning")
            return
        meta = game.get("meta") or {}
        title = meta.get("title", "Unknown")
        pid, gid = get_game_ids(game)
        if not pid or not gid:
            self.notify(f"Game «{title}» has no master_playerid/master_id; synchronise games again (s)", severity="error")
            return
        self.push_screen(AwardsScreen(title, platform or "?", pid, gid, self._output_dir(), self._cache_dir()))

    def action_export_awards_history(self) -> None:
        if not self.games_data:
            self.notify("No data to export", severity="error")
            return
        self.notify("Collecting achievement history for all games… (this may take several minutes)", severity="information")
        self.export_awards_history_data()

    def action_sync_all_awards(self) -> None:
        """Синк відсутніх ачивок (пропускає кеш). Повторний виклик під час роботи — зупинка."""
        if self._awards_sync_running:
            self._awards_sync_stop = True
            self.notify("⏹ Stopping achievement synchronisation…", severity="warning")
            return
        if not self.games_data:
            self.notify("No data; synchronise games first (s)", severity="error")
            return
        self.notify("⏳ Syncing achievements for all games (skipping cached games)… press w again to stop", severity="information")
        self.sync_all_awards_data(force=False)

    def action_sync_all_awards_force(self) -> None:
        """Пересинк ВСІХ ачивок заново (ігнорує кеш). Повторний виклик — зупинка."""
        if self._awards_sync_running:
            self._awards_sync_stop = True
            self.notify("⏹ Stopping achievement synchronisation…", severity="warning")
            return
        if not self.games_data:
            self.notify("No data; synchronise games first (s)", severity="error")
            return
        self.notify("⏳ Resyncing ALL achievements (ignoring cache)… this takes time. Press W again to stop", severity="information")
        self.sync_all_awards_data(force=True)

    def _show_awards_progress(self, done: int, total: int, fetched: int, skipped: int, errors: int) -> None:
        try:
            file_info = self.latest_json_path.name if self.latest_json_path else "—"
            self.query_one("#statusbar", Static).update(
                f" 📁 {file_info}   |   ⏳ achievements: {done}/{total} (fetched: {fetched}, cache: {skipped}, errors: {errors})"
            )
        except Exception:
            pass

    @work(thread=True)
    def sync_all_awards_data(self, force: bool = False) -> None:
        cache_dir = self._cache_dir()
        self._awards_sync_running = True
        self._awards_sync_stop = False
        games = list(self.games_data)
        total = len(games)
        done = 0
        fetched = 0
        skipped = 0
        errors = 0
        total_awards = 0
        for game in games:
            if self._awards_sync_stop:
                break
            if not isinstance(game, dict):
                done += 1
                continue
            meta = game.get("meta") or {}
            pid, gid = get_game_ids(game)
            if not pid or not gid:
                errors += 1
                done += 1
                continue
            cp = awards_cache_path(cache_dir, pid, gid)
            if not force:
                try:
                    if cp.exists():
                        with open(cp, "r", encoding="utf-8") as f:
                            cached = json.load(f)
                        if isinstance(cached, list):
                            skipped += 1
                            total_awards += len(cached)
                            done += 1
                            if done % 5 == 0:
                                self.call_from_thread(
                                    self._show_awards_progress, done, total, fetched, skipped, errors
                                )
                            continue
                except Exception:
                    pass  # битий кеш — перетягнемо з API
            try:
                lst = fetch_earned_awards(pid, gid)
                try:
                    with open(cp, "w", encoding="utf-8") as f:
                        json.dump(lst, f, ensure_ascii=False, indent=2)
                except Exception:
                    pass
                fetched += 1
                total_awards += len(lst)
                time.sleep(0.25)  # не душити API
            except Exception:
                errors += 1
            done += 1
            if done % 5 == 0 or done == total:
                self.call_from_thread(
                    self._show_awards_progress, done, total, fetched, skipped, errors
                )
        self._awards_sync_running = False
        stopped = self._awards_sync_stop
        self._awards_sync_stop = False
        # повернути статусбар до звичайного вигляду
        try:
            loaded = len(games)
            self.call_from_thread(self._update_statusbar, loaded, "")
        except Exception:
            pass
        if stopped:
            self.call_from_thread(
                self.notify,
                f"⏹ Stopped: {done}/{total} games (fetched: {fetched}, cached: {skipped}, errors: {errors})",
                severity="warning",
            )
        else:
            self.call_from_thread(
                self.notify,
                f"✅ Achievement sync complete: {done}/{total} games, {total_awards} achievements "
                f"(fetched: {fetched}, cached: {skipped}, errors: {errors})"
                + ("; older downloads without master_id require synchronising games again (s)" if errors else ""),
                severity="information",
            )

    @work(thread=True)
    def export_awards_history_data(self) -> None:
        out_dir = self._output_dir()
        cache_dir = self._cache_dir()
        player_id = (self.latest_json_path.stem.split("_")[-1]
                     if self.latest_json_path else "export")
        all_rows: list[dict] = []
        total_games = len(self.games_data)
        done = 0
        errors = 0
        for game in list(self.games_data):
            if not isinstance(game, dict):
                continue
            meta = game.get("meta") or {}
            title = meta.get("title", "Unknown")
            platform = classify_platform(meta.get("platforms") or [])
            pid, gid = get_game_ids(game)
            if not pid or not gid:
                errors += 1
                continue
            lst: list[dict] | None = None
            # 1) кеш
            try:
                cp = awards_cache_path(cache_dir, pid, gid)
                if cp.exists():
                    with open(cp, "r", encoding="utf-8") as f:
                        cached = json.load(f)
                    if isinstance(cached, list):
                        lst = cached
            except Exception:
                lst = None
            # 2) API, якщо кешу нема
            if lst is None:
                try:
                    lst = fetch_earned_awards(pid, gid)
                    try:
                        with open(awards_cache_path(cache_dir, pid, gid), "w", encoding="utf-8") as f:
                            json.dump(lst, f, ensure_ascii=False, indent=2)
                    except Exception:
                        pass
                    time.sleep(0.3)  # не душити API
                except Exception:
                    errors += 1
                    continue
            for aw in lst or []:
                if not isinstance(aw, dict):
                    continue
                name, earned_utc, earned_site, aid = build_award_row(aw)
                all_rows.append({
                    "Game": title,
                    "Platform": platform,
                    "Achievement": name,
                    "Earned (UTC)": earned_utc,
                    "Earned (site)": earned_site,
                    "Award ID": aid,
                })
            done += 1
            if done % 10 == 0:
                self.call_from_thread(
                    self.notify, f"…{done}/{total_games} games, {len(all_rows)} achievements", severity="information"
                )

        if not all_rows:
            self.call_from_thread(self.notify, "No achievement history collected (check synchronisation)", severity="error")
            return
        # сортуємо за датою — найновіші зверху
        all_rows.sort(key=lambda r: r.get("Earned (UTC)", ""), reverse=True)
        base = out_dir / f"exophase_awards_history_{player_id}"
        try:
            # Excel
            xlsx_path = base.with_suffix(".xlsx")
            df = pd.DataFrame(all_rows, columns=list(HISTORY_COLUMNS))
            with report_file(xlsx_path) as working_path:
                with pd.ExcelWriter(working_path, engine="openpyxl") as writer:
                    df.to_excel(writer, sheet_name="History", index=False)
                    ws = writer.sheets["History"]
                    for cell in ws[1]:
                        cell.fill = PatternFill("solid", fgColor="1F4E79")
                        cell.font = Font(color="FFFFFF", bold=True)
                        cell.alignment = Alignment(horizontal="center")
                    for col_idx, col in enumerate(ws.columns, 1):
                        width = max((len(str(c.value)) for c in col if c.value), default=8)
                        ws.column_dimensions[get_column_letter(col_idx)].width = min(width + 2, 60)
            # CSV
            csv_path = base.with_suffix(".csv")
            with report_file(csv_path) as working_path:
                with open(working_path, "w", newline="", encoding="utf-8-sig") as f:
                    writer = csv.DictWriter(f, fieldnames=list(HISTORY_COLUMNS))
                    writer.writeheader()
                    writer.writerows(all_rows)
            self.call_from_thread(
                self.notify,
                f"✅ History: {len(all_rows)} achievements from {done} games → {xlsx_path.name}, {csv_path.name}"
                + (f" (skipped: {errors})" if errors else ""),
                severity="information",
            )
        except Exception as e:
            self.call_from_thread(self.notify, f"History export error: {e}", severity="error")

    def action_delete_game(self) -> None:
        dt = self._active_datatable()
        if dt is None:
            return
        try:
            row_key    = dt.coordinate_to_cell_key(dt.cursor_coordinate).row_key
            row_values = dt.get_row(row_key)
            target_title = str(row_values[0])
        except Exception:
            self.notify("No row selected", severity="warning")
            return

        for i, game in enumerate(self.games_data):
            meta = game.get("meta") or {}
            if meta.get("title", "") == target_title:
                del self.games_data[i]
                break

        self._save_games_data()
        self.load_data(self._filter_query)
        self.notify(f"Deleted: {target_title}", severity="information")

    def action_sync(self) -> None:
        def on_result(player_id: str | None) -> None:
            if player_id:
                self.notify(f"Synchronising player «{player_id}»…", severity="information")
                self.fetch_api_data(player_id)

        self.push_screen(SyncModal(), on_result)

    def action_export(self) -> None:
        self.notify("Creating Excel file…", severity="information")
        self.export_excel_data()

    def action_export_csv(self) -> None:
        self.notify("Creating CSV file…", severity="information")
        self.export_csv_data()

    @work(thread=True)
    def fetch_api_data(self, player_id: str) -> None:
        cache_dir = self._cache_dir()

        all_games: list[dict] = []
        for page_num in range(1, 9999):
            url = (
                f"https://api.exophase.com/public/player/{player_id}/games"
                f"?page={page_num}&environment=&sort=1&showHidden=0"
            )
            try:
                resp = requests.get(url, headers=API_HEADERS, timeout=10)
                resp.raise_for_status()
                data = resp.json()
            except Exception as e:
                self.call_from_thread(self.notify, f"API error (page {page_num}): {e}", severity="error")
                break

            if not data.get("success"):
                break
            games = data.get("games") or []
            if not games:
                break
            all_games.extend(games)

        if not all_games:
            self.call_from_thread(self.notify, "No data received from API", severity="warning")
            return

        json_path = cache_dir / f"all_games_{player_id}.json"
        try:
            with open(json_path, "w", encoding="utf-8") as f:
                json.dump(all_games, f, ensure_ascii=False, indent=4)
            self.call_from_thread(
                self.notify, f"✅ Synchronised {len(all_games)} games → {json_path.name}", severity="information"
            )
            self.call_from_thread(self.load_data, self._filter_query)
        except Exception as e:
            self.call_from_thread(self.notify, f"Save error: {e}", severity="error")

    @work(thread=True)
    def export_excel_data(self) -> None:
        out_dir = self._output_dir()
        if not self.games_data:
            self.call_from_thread(self.notify, "No data to export", severity="error")
            return

        sheets: dict[str, list[dict]] = {k: [] for k in ("PlayStation", "Xbox", "Steam", "Other")}

        for game in self.games_data:
            meta      = game.get("meta") or {}
            platforms = meta.get("platforms") or []
            platform  = classify_platform(platforms)
            row_tuple = build_row(game, platform)
            col_names = TABLE_COLUMNS[PLATFORM_TO_DT[platform].lstrip("#")]
            sheets[platform].append(dict(zip(col_names, row_tuple)))

        player_id   = (self.latest_json_path.stem.split("_")[-1]
                       if self.latest_json_path else "export")
        excel_path  = out_dir / f"exophase_games_{player_id}.xlsx"

        HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")
        HEADER_FONT  = Font(color="FFFFFF", bold=True)
        HEADER_ALIGN = Alignment(horizontal="center")

        try:
            with report_file(excel_path) as working_path:
                with pd.ExcelWriter(working_path, engine="openpyxl") as writer:
                    for sheet_name, rows in sheets.items():
                        if not rows:
                            continue
                        df = pd.DataFrame(rows)
                        df.to_excel(writer, sheet_name=sheet_name, index=False)
                        ws = writer.sheets[sheet_name]

                        for cell in ws[1]:
                            cell.fill      = HEADER_FILL
                            cell.font      = HEADER_FONT
                            cell.alignment = HEADER_ALIGN

                        for col_idx, col in enumerate(ws.columns, 1):
                            width = max((len(str(c.value)) for c in col if c.value), default=8)
                            ws.column_dimensions[get_column_letter(col_idx)].width = width + 2

            self.call_from_thread(
                self.notify, f"✅ Excel → {excel_path.name}", severity="information"
            )
        except Exception as e:
            self.call_from_thread(self.notify, f"Excel export error: {e}", severity="error")

    @work(thread=True)
    def export_csv_data(self) -> None:
        if not self.games_data:
            self.call_from_thread(self.notify, "No data to export as CSV", severity="error")
            return

        out_dir   = self._output_dir()
        tabs      = self.query_one(TabbedContent)
        active_id = tabs.active
        platform  = active_id.replace("tp-", "").capitalize()
        if platform == "Ps":
            platform = "PlayStation"

        rows: list[dict] = []
        for game in self.games_data:
            meta  = game.get("meta") or {}
            plat  = classify_platform(meta.get("platforms") or [])
            if plat.lower() != platform.lower():
                continue
            row_tuple = build_row(game, plat)
            col_names = TABLE_COLUMNS[PLATFORM_TO_DT[plat].lstrip("#")]
            rows.append(dict(zip(col_names, row_tuple)))

        if not rows:
            self.call_from_thread(self.notify, f"No data for the {platform} tab", severity="warning")
            return

        player_id = (self.latest_json_path.stem.split("_")[-1]
                     if self.latest_json_path else "export")
        csv_path  = out_dir / f"exophase_{platform.lower()}_{player_id}.csv"

        try:
            with report_file(csv_path) as working_path:
                with open(working_path, "w", newline="", encoding="utf-8-sig") as f:
                    writer = csv.DictWriter(f, fieldnames=rows[0].keys())
                    writer.writeheader()
                    writer.writerows(rows)
            self.call_from_thread(
                self.notify, f"✅ CSV → {csv_path.name}", severity="information"
            )
        except Exception as e:
            self.call_from_thread(self.notify, f"CSV export error: {e}", severity="error")


if __name__ == "__main__":
    SELFexophase().run()